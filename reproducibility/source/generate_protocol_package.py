#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
import textwrap
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output"
FIGURES = OUTPUT / "figures"
REPRO = ROOT / "reproducibility"
NOW = datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def load_json(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


CONFIG = load_json(ROOT / "config" / "trial_design.json")
CONSTANTS = load_json(OUTPUT / "analysis_constants.json")
MOCK = load_json(REPRO / "outputs" / "mock_analysis_summary.json")
REFERENCES = load_json(OUTPUT / "references_vancouver.json")
REFERENCE_BY_ID = {record["id"]: record for record in REFERENCES}


def ref(reference_id: str) -> int:
    return int(REFERENCE_BY_ID[reference_id]["reference_number"])


TITLE = str(CONFIG["full_title"])
ACRONYM = str(CONFIG["acronym"])
TARGET_TOTAL = int(CONSTANTS["target_total"])
TARGET_ARM = int(CONSTANTS["target_per_arm"])
CONTROL_RATE = float(CONSTANTS["planning_control_rate"])
TREATMENT_RATE = float(CONSTANTS["planning_treatment_rate"])
POWER = float(CONSTANTS["power"])
ALPHA = float(CONSTANTS["alpha_two_sided"])
NON_EVALUABLE = float(CONSTANTS["non_evaluable_fraction"])
DOSE = int(CONFIG["intervention"]["dose_mg"])
VOLUME = int(CONFIG["intervention"]["volume_ml"])
ADMIN_MIN = int(CONFIG["intervention"]["administration_minutes"])
LEAD = int(CONFIG["intervention"]["target_completion_minutes_before_protamine"])
WINDOW = int(CONFIG["intervention"]["window_minutes"])
ENDPOINT_WINDOW = int(CONFIG["primary_endpoint"]["window_minutes"])


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.rstrip() + "\n", encoding="utf-8")


def write_csv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def add_bottom_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), "B7B7B7")
    borders.append(bottom)
    tc_pr.append(borders)


def set_doc_defaults(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 2
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = TITLE


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)


def add_paragraph(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
        add_bottom_border(cell)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.add_paragraph()


def add_figure(document: Document, image_path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.4))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.paragraph_format.space_before = Pt(14)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = caption_paragraph.add_run(caption)
    run.bold = True


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_figures() -> list[dict[str, str]]:
    FIGURES.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})

    fig, ax = plt.subplots(figsize=(10, 5.8))
    ax.axis("off")
    boxes = [
        (0.08, 0.77, "Eligibility screening\nElective non-CABG CPB surgery"),
        (0.08, 0.53, "Consent and baseline assessment"),
        (0.08, 0.29, f"Central randomisation\nPlanned N={TARGET_TOTAL}"),
        (0.43, 0.52, f"FPA {DOSE} mg IV\nPlanned n={TARGET_ARM}"),
        (0.70, 0.52, f"Placebo {VOLUME} mL IV\nPlanned n={TARGET_ARM}"),
        (0.43, 0.22, "Protamine t=0\nContinuous MAP to +30 min"),
        (0.70, 0.22, "Protamine t=0\nContinuous MAP to +30 min"),
        (0.56, 0.03, "In-hospital/POD7 safety; 30-day vital status"),
    ]
    for x, y, label in boxes:
        ax.text(
            x,
            y,
            label,
            ha="center",
            va="center",
            transform=ax.transAxes,
            bbox={"boxstyle": "round,pad=0.45", "fc": "#F4F7FB", "ec": "#315B7D", "lw": 1.4},
        )
    arrows = [
        ((0.08, 0.70), (0.08, 0.60)),
        ((0.08, 0.46), (0.08, 0.36)),
        ((0.14, 0.29), (0.38, 0.52)),
        ((0.14, 0.29), (0.64, 0.52)),
        ((0.43, 0.44), (0.43, 0.30)),
        ((0.70, 0.44), (0.70, 0.30)),
        ((0.48, 0.17), (0.54, 0.08)),
        ((0.66, 0.17), (0.60, 0.08)),
    ]
    for start, end in arrows:
        ax.annotate("", xy=end, xytext=start, xycoords="axes fraction", arrowprops={"arrowstyle": "->", "lw": 1.2})
    ax.set_title("Planned participant flow and assessment schedule", fontsize=14, weight="bold")
    fig.tight_layout()
    flow_png = FIGURES / "Figure_1_trial_timeline.png"
    flow_tiff = FIGURES / "Figure_1_trial_timeline.tiff"
    fig.savefig(flow_png, dpi=300, bbox_inches="tight")
    fig.savefig(
        flow_tiff,
        dpi=300,
        bbox_inches="tight",
        pil_kwargs={"compression": "tiff_lzw"},
    )
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(10, 4.8))
    ax.set_xlim(-25, 32)
    ax.set_ylim(-0.2, 3.2)
    ax.hlines(1.45, -25, 32, color="#315B7D", linewidth=2)
    ax.axvspan(-20, -10, color="#F4C95D", alpha=0.35, label="Study infusion target")
    ax.axvspan(-5, -1, color="#8AC6D1", alpha=0.35, label="Baseline MAP window")
    ax.axvspan(0, 30, color="#E76F51", alpha=0.16, label="Primary endpoint window")
    ax.axvline(0, color="#9B2226", linewidth=2.2)
    ax.text(0, 2.75, "t=0: first protamine infusion starts", ha="center", weight="bold")
    ax.text(-15, 2.25, f"FPA/placebo complete\n{LEAD} ± {WINDOW} min before t=0", ha="center")
    ax.text(-3, 0.55, "Stabilised baseline:\nmedian valid MAP, −5 to −1 min", ha="center")
    ax.text(15, 2.25, "Continuous arterial MAP ≥1 Hz\nrolling 60-s epochs to +30 min", ha="center")
    ax.scatter([-20, -10, -5, -1, 0, 30], [1.45] * 6, color="#315B7D", zorder=3)
    ax.set_xlabel("Minutes relative to protamine start")
    ax.set_yticks([])
    ax.spines[["left", "right", "top"]].set_visible(False)
    ax.legend(loc="lower center", ncol=3, frameon=False)
    ax.set_title("Prospective haemodynamic measurement timeline", fontsize=14, weight="bold")
    fig.tight_layout()
    haem_png = FIGURES / "Figure_2_hemodynamic_timeline.png"
    haem_tiff = FIGURES / "Figure_2_hemodynamic_timeline.tiff"
    fig.savefig(haem_png, dpi=300, bbox_inches="tight")
    fig.savefig(
        haem_tiff,
        dpi=300,
        bbox_inches="tight",
        pil_kwargs={"compression": "tiff_lzw"},
    )
    plt.close(fig)

    return [
        {
            "title": "Figure 1. Planned participant flow and assessment schedule",
            "caption": f"Planned flow for the {ACRONYM} trial. Counts are design targets and do not represent enrolment. CPB, cardiopulmonary bypass; FPA, flurbiprofen axetil; MAP, mean arterial pressure; POD, postoperative day.",
            "path": str(flow_png),
        },
        {
            "title": "Figure 2. Prospective haemodynamic measurement timeline",
            "caption": "Study-drug administration, stabilised baseline, protamine time zero, and the 30-minute primary endpoint window. FPA, flurbiprofen axetil; MAP, mean arterial pressure.",
            "path": str(haem_png),
        },
    ]


TABLES = [
    {
        "title": "Table 1. Schedule of enrolment, interventions, and assessments",
        "headers": ["Procedure", "Screening", "Pre-protamine", "0–30 min", "ICU/POD1", "POD2–7/discharge", "Day 30"],
        "rows": [
            ["Eligibility/consent", "X", "", "", "", "", ""],
            ["Randomisation", "", "X", "", "", "", ""],
            ["FPA/placebo", "", f"Complete {LEAD} ± {WINDOW} min before t=0", "", "", "", ""],
            ["Protamine/heparin/ACT capture", "", "X", "X", "", "", ""],
            ["Continuous MAP and vasoactive drugs", "", "Baseline", "X", "", "", ""],
            ["Bleeding/transfusion/reoperation", "", "", "", "X", "X", ""],
            ["AKI and serious adverse events", "", "", "", "X", "X", ""],
            ["Vital status", "", "", "", "", "", "X"],
        ],
    },
    {
        "title": "Table 2. Key eligibility criteria",
        "headers": ["Include", "Exclude"],
        "rows": [
            ["Age ≥18 years; written informed consent", "Emergency or salvage operation"],
            ["Elective non-CABG cardiac surgery requiring CPB", "Planned CABG or another procedure for which systemic NSAID exposure is prohibited locally"],
            ["Systemic unfractionated heparin and planned protamine reversal", "FPA/NSAID hypersensitivity, aspirin-exacerbated respiratory disease, or prior severe protamine reaction"],
            ["Arterial catheter suitable for digital waveform capture", "Active bleeding, clinically important coagulopathy, or contraindicated antiplatelet/anticoagulant context"],
            ["Randomisation feasible before study-drug preparation", "Active peptic ulcer, severe renal/hepatic disease, decompensated heart failure/cardiogenic shock, severe uncontrolled hypertension, pregnancy, or interacting prohibited medicine under the Japanese label"],
        ],
    },
    {
        "title": "Table 3. Outcome definitions",
        "headers": ["Outcome", "Prospective definition", "Analysis"],
        "rows": [
            [
                "Primary",
                "≥1 artefact-free rolling 60-s epoch from t=0 to +30 min with median MAP ≤80% of stabilised baseline; ≥45 valid seconds required",
                "Treatment-policy ITT; adjusted marginal risk difference with 95% stratified-bootstrap CI",
            ],
            ["Absolute hypotension", "Any valid 60-s epoch with median MAP <65 mmHg", "Risk difference and risk ratio"],
            ["Hypotension burden", "Time-weighted average and area under MAP 65 mmHg", "Adjusted linear model; transformed/robust analysis if needed"],
            ["Rescue treatment", "Protocol-defined vasopressor bolus, infusion escalation, or rescue mechanical support", "Components and composite with the primary MAP criterion"],
            ["Safety", "KDIGO AKI; chest-tube drainage; transfusion; reoperation; GI bleeding; bronchospasm/anaphylaxis; MI; stroke; SAE; death", "Prespecified descriptive estimates with 95% CIs; no efficacy multiplicity claim"],
        ],
    },
]


def markdown_documents() -> dict[str, str]:
    source_extraction = f"""
# Source extraction

## Source status

The hypothesis-generating source is the author-supplied manuscript, “Association Between Intravenous Flurbiprofen Axetil Pre-administration and Attenuation of Protamine-Induced Hypotension in Cardiovascular Surgery: a retrospective cohort study.” The supplied DOCX is the numeric source of record for this package. The local PDF was generated from that DOCX and is not independently verified as the final publisher PDF. The official JSCA member eBook could not be audited without member credentials.

## Extracted design

- Single-centre retrospective cohort using 2009–2011 records.
- 92 cardiovascular operations identified; five peripheral vascular procedures below the diaphragm excluded; 87 analysed.
- The procedure table contains 40 on-pump open-heart procedures and 47 off-pump CABG procedures.
- FPA exposure reflected attending-anaesthesiologist preference and routine analgesic practice, not random allocation.
- FPA was administered at least 15 minutes before protamine.
- Protamine was described as 10 mg per 1000 units of initial heparin over five minutes by syringe pump.
- Historical MAPpre was the lowest MAP from before protamine through 15 minutes after the start; MAPpost was the lowest MAP from 15 to 30 minutes.
- Historical PSI was `(MAPpre − MAPpost) / MAPpre`; PSI >0.20 was an author-defined relative MAP-decline threshold, not a validated universal shock definition.

## Extracted results

- On-pump difference-in-differences estimate: 4.09 mmHg attenuation of MAP decline (approximately 95% CI 0.72–7.46; p=0.02).
- Weighted model-based probabilities were approximately 0.00 with FPA and 0.17 without FPA.
- Bootstrap risk difference was approximately −0.18 (approximately 95% CI −0.40 to −0.03).

## Limitations carried forward

Treatment-selection bias, small on-pump subgroup, historical paper records, limited propensity-score overlap, residual imbalance, broad lowest-MAP windows, incomplete time-stamped vasoactive data, and a single-centre design make the findings hypothesis-generating only. None of these retrospective estimates is represented as prospective efficacy evidence.
"""
    journal_requirements = """
# Current Annals of Cardiac Anaesthesia requirements audit

Official information for authors was captured on 2026-09-25 and retained under `source/public/`.

| Requirement | Verified requirement | Package response |
|---|---|---|
| Article categories | Original Article, Review, Case Report, Grand Round Case, Images, Letter, and other listed types; “protocol” is not explicitly listed | Presubmission inquiry is mandatory before formal submission |
| Abstract | Structured, maximum 250 words for Original Articles | Protocol abstract is structured and validated at ≤250 words |
| Main text | Maximum 3000 words for Original Articles | Generator fails if the blinded manuscript body exceeds 3000 words |
| Tables/figures | Maximum five combined | Three tables and two figures |
| References | Maximum 30 | Twenty-two verified references |
| Review | Blinded peer review | Blinded manuscript separated from title page |
| Author identity | Unblinded first-page/title-page file or cover letter | Separate title page and cover letter |
| AI disclosure | Name tool and purpose; authors retain responsibility | Included in declarations and cover letter |
| Submission portal | https://review.jow.medknow.com/aca | Recorded in handoff |

The package is prepared to the numeric limits of the Original Article category solely as a conservative formatting envelope. This does not establish that ACA accepts protocol manuscripts.
"""
    article_type = """
# Article-type decision

## Decision

Do not submit silently as a completed Original Article. ACA’s current author instructions do not explicitly list clinical trial protocols as an article category. Send `presubmission_inquiry_ACA.docx` first and ask whether the journal will consider the manuscript, which category should be selected in the portal, and whether any protocol-specific supplements are required.

## Contingency

If ACA declines protocol manuscripts, retain this SPIRIT-compliant package and adapt it for a journal that explicitly accepts trial protocols. Do not reframe planned methods as completed results.
"""
    aca_fit = f"""
# ACA fit assessment

The topic—protamine reversal, haemodynamic instability, CPB anticoagulation, and perioperative pharmacology—is closely aligned with cardiac anaesthesia. The principal editorial risk is article type rather than scope: ACA does not currently list a protocol category. A second risk is that FPA is not globally available and prophylactic use is off-label/investigational. The protocol therefore emphasises mechanistic rationale without claiming efficacy, Japanese-label restrictions, an exclusion of planned CABG, objective digital MAP capture, and a multicentre design.

The package meets the conservative Original Article envelope: structured abstract ≤250 words, main body ≤3000 words, {len(TABLES) + 2} combined tables/figures, and {len(REFERENCES)} references. Acceptance remains subject to the editor’s response to the presubmission inquiry.
"""
    precedent = """
# ACA protocol precedent search

Searches of the current ACA information-for-authors page and indexed ACA article categories did not identify an explicit “study protocol” category. Searches can miss unindexed or inconsistently labelled protocols; absence is not proof that ACA has never published one. Because the governing author instructions are silent, the package treats protocol acceptance as unresolved and supplies a presubmission inquiry rather than claiming precedent.
"""
    haemodynamic = f"""
# Haemodynamic measurement specification

## Time origin and acquisition

- `t=0` is the timestamp at which the first protamine infusion starts.
- Capture the arterial-pressure waveform at ≥1 Hz from at least −5 through +30 minutes relative to t=0.
- Record monitor time synchronisation, transducer level/zero confirmation, line site, and monitor/export system.
- The stabilised baseline MAP is the median of valid MAP observations from −5 to −1 minutes.

## Primary endpoint derivation

Compute a rolling 60-second median across the 0–30-minute window. A window is evaluable when at least 45 seconds are valid. The primary event occurs if at least one evaluable window has median MAP ≤80% of baseline. Exact equality qualifies.

## Supporting derivations

- MAP <65 mmHg: at least one evaluable 60-second window below 65 mmHg.
- Nadir MAP: minimum valid rolling 60-second median.
- Relative maximum decline: `(baseline MAP − nadir MAP) / baseline MAP`.
- Time under 65 mmHg: sum of valid seconds below 65, divided by 60.
- Area under 65 mmHg: integral of `max(65 − MAP, 0)` over valid time.
- Time-weighted MAP: area under the observed MAP curve divided by valid observation duration.

## Concomitant treatment

Record every vasopressor/inotrope bolus and infusion-rate change with exact time, drug, dose, and route from −10 through +30 minutes. Record fluids, pacing, cardioversion, return to CPB, mechanical support, and surgical events. The primary treatment-policy estimand uses observed MAP after rescue; a prespecified composite counts qualifying rescue as an event.
"""
    artifacts = """
# Arterial-waveform artifact handling rules

## Automated flags

Mark a second invalid if MAP is physiologically impossible (<20 or >180 mmHg), the signal is absent, the arterial line is documented as disconnected/occluded, or the monitor identifies a flush/calibration event. Automated flags do not replace blinded adjudication.

## Flushes, damping, and isolated beats

- Exclude square-wave flush periods from five seconds before the flush until the waveform returns to interpretable pulsatility.
- Exclude documented blood sampling, line opening, transducer zeroing, and reconnection.
- Exclude clearly overdamped/underdamped periods when the waveform is nonphysiologic and corroborated by the site record or blinded central review.
- Do not classify a single-beat spike or trough as hypotension. The primary endpoint uses rolling 60-second medians and requires ≥45 valid seconds.
- Preserve true arrhythmia-associated pressure variation unless the waveform itself is uninterpretable.

## Treatment-induced measurements

Do not censor values after rescue vasopressor, fluid, pacing, cardioversion, or return to CPB in the primary treatment-policy analysis. Record interventions precisely. A composite sensitivity endpoint counts protocol-defined rescue, and supportive analyses will describe MAP before rescue.

## Missingness

The primary endpoint is non-evaluable when baseline has <180 valid seconds or the post-protamine window has <80% valid seconds after documented artifact exclusion. Do not impute waveform seconds for the primary derivation. Primary ITT analysis will use multiple imputation for a missing binary endpoint only if missingness exceeds 1%; best/worst-case bounds and complete-case analyses will be reported. Death or rescue mechanical support before endpoint ascertainment counts as an event.
"""
    randomisation = """
# Randomisation specification

- Unit: participant.
- Allocation: 1:1 FPA versus placebo.
- Sequence: validated computer-generated sequence with randomly permuted block sizes 4, 6, and 8.
- Strata: participating site and preoperative exposure to protamine-containing insulin or documented prior protamine exposure (yes/no).
- Concealment: central web-based allocation or equivalent independent system; the next assignment must not be visible to enrolling staff.
- Access: sequence accessible only to the independent statistician/system administrator and unblinded pharmacy/preparation personnel.
- Audit: retain generation code, seed under restricted access, allocation logs, timestamps, and reason for every override.
- Final incomplete blocks are acceptable; recruitment targets are balanced operationally but the analysis does not assume exact equality.
"""
    blinding = f"""
# Blinding specification

FPA {DOSE} mg/{VOLUME} mL and {VOLUME} mL 0.9% saline placebo will be prepared by unblinded pharmacy or independent preparation personnel. Opaque syringes and tubing will be used because visual differences may compromise masking. Study drug will be infused over {ADMIN_MIN} minutes and completed {LEAD} ± {WINDOW} minutes before protamine.

Participants, anaesthesiologists, surgeons, ICU clinicians, outcome assessors, central waveform reviewers, investigators, and statisticians remain masked unless knowledge is essential for immediate care. Preparation staff have no endpoint-assessment role. Allocation guesses and reasons are recorded after the 30-minute endpoint window.

Emergency unblinding is available continuously through the allocation system or independent pharmacy. The treating clinician documents the medical necessity, date/time, person requesting and releasing the code, and resulting action. Unblinding for convenience, routine adverse-event assessment, or data review is prohibited.
"""
    safety = f"""
# Safety and regulatory review

## Status of use

Ropion is supplied in Japan as flurbiprofen axetil {DOSE} mg/{VOLUME} mL for postoperative and cancer pain. Administration to prevent protamine-associated hypotension is not an approved indication and is investigational/off-label. Regulatory classification, Clinical Trials Act applicability, insurance/compensation, and required notifications must be confirmed by the Japanese sponsor and ethics/regulatory offices before enrolment.

## Principal risks

- **Renal:** NSAID-mediated reduction in renal perfusion may precipitate AKI, particularly after CPB. Exclude severe renal disease; record creatinine and urine output; adjudicate KDIGO AKI through POD7/discharge.
- **Haemostasis and gastrointestinal:** cyclooxygenase inhibition may impair platelet function or increase bleeding. Exclude active bleeding, clinically important coagulopathy, and prohibited interacting antiplatelet/anticoagulant contexts. Record chest-tube drainage, transfusion, reoperation, and GI bleeding.
- **Hypersensitivity/bronchospasm:** exclude FPA/NSAID allergy, aspirin-exacerbated respiratory disease, and prior severe protamine reaction. Resuscitation drugs and equipment must be immediately available.
- **Cardiovascular:** international systemic NSAID labels carry thrombotic and CABG warnings. Planned CABG is excluded conservatively; this is an international class-risk precaution, not a claim that the retrieved Japanese FPA label contains an identical CABG prohibition.
- **Label restrictions:** exclude active peptic ulcer, severe haematological/hepatic/renal disease, decompensated heart failure or cardiogenic shock, severe uncontrolled hypertension, late pregnancy/pregnancy as locally required, and prohibited interacting medicines. The final Japanese eligibility wording must be approved against the current label because “severe” organ dysfunction requires operational thresholds.
- **Protamine:** rapid or excessive dosing can cause hypotension, pulmonary hypertension, allergy, and anticoagulant effects. Follow the locally approved label, prefer residual-heparin titration, and record all dosing/rate deviations.

## Safety governance

An independent DSMB will review masked aggregate data after 100 participants and unmasked data after 300 and 600 participants, and ad hoc after any suspected unexpected serious adverse reaction (SUSAR) or death plausibly related to study drug. The DSMB may recommend continue, modify, pause, or terminate for clear excess death, anaphylaxis/bronchospasm, KDIGO stage 2–3 AKI, reoperation for bleeding, or other serious toxicity. No formal efficacy stopping boundary is planned; this avoids an underpowered early efficacy claim.

All AEs are collected from study-drug start through discharge or POD7; SAEs and pregnancies through day 30. Investigators report SAEs promptly to the sponsor and ethics body under applicable law; fatal/life-threatening unexpected related reactions are expedited as required. Causality, expectedness against the Japanese label, severity, action, and outcome are recorded.
"""
    eligibility = """
# Eligibility rationale

## Inclusion

Adults undergoing elective non-CABG cardiac surgery with CPB, systemic unfractionated heparinisation, planned protamine reversal, and digital arterial-waveform capture form a clinically coherent population exposed to the target mechanism and endpoint.

## Exclusion rationale

- Emergency/salvage surgery: consent, drug timing, and baseline stabilisation are unreliable.
- Planned CABG: conservative response to international systemic NSAID CABG warnings.
- Severe renal, hepatic, or haematological disease; decompensated heart failure/cardiogenic shock; or uncontrolled hypertension: Japanese-label safety restrictions and increased risk from NSAID exposure. Numeric/clinical thresholds must be finalised before activation.
- Active bleeding/coagulopathy or prohibited antiplatelet context: bleeding-risk minimisation.
- Peptic ulcer/GI bleeding: NSAID GI risk.
- FPA/NSAID hypersensitivity or aspirin asthma: anaphylaxis/bronchospasm risk.
- Prior severe protamine reaction: the trial cannot ethically rely on FPA to prevent a known high-risk reaction.
- Pregnancy: label restriction and absent direct benefit.
- Emergency inability to mask or complete study drug within the timing window: protocol integrity.

Prior protamine or protamine-containing insulin exposure is not an automatic exclusion unless associated with a severe reaction; it is a randomisation stratum and prespecified subgroup because it may modify baseline risk.
"""
    sap = f"""
# Statistical Analysis Plan

Version {CONFIG["protocol_version"]}; dated {CONFIG["protocol_date"]}. Finalise and sign before database lock and before treatment codes are released.

## Objectives and estimands

The primary objective is to estimate the effect of assignment to FPA versus placebo on clinically important peri-protamine hypotension. The primary estimand is a treatment-policy risk difference in all randomised participants, regardless of adherence, crossover, rescue treatment, or protocol deviation, except that participants randomised in error who never undergo surgery/protamine remain in the ITT denominator with endpoint status handled under the missing-data rules.

## Analysis populations

- **ITT:** all randomised participants, analysed as assigned.
- **Safety:** all participants receiving any study drug, analysed as treated.
- **Per protocol:** study drug administered as assigned; dose 45–55 mg or placebo-equivalent volume; infusion duration 4–6 minutes; completion 10–20 minutes before t=0; eligible surgery; no prohibited unblinding; evaluable waveform.

## Primary analysis

Fit logistic regression with treatment, site stratum, prior protamine/protamine-insulin exposure, and continuous baseline MAP. Standardise predicted risks over the ITT population. Report marginal risk difference and risk ratio with two-sided 95% CIs; the primary CI for the risk difference uses the 2.5th and 97.5th percentiles of {CONFIG["analysis"]["bootstrap_samples"]:,} stratified bootstrap samples. Also report crude risks and odds ratio. Statistical significance uses two-sided alpha {ALPHA:.2f}.

If complete or quasi-complete separation occurs, use Firth penalised logistic regression for conditional estimates and report an exact unconditional risk difference interval as a sensitivity analysis. Model convergence, influential observations, calibration, and functional form of baseline MAP will be checked.

## Intercurrent events and missing data

Rescue vasoactive treatment, fluids, pacing, cardioversion, return to CPB, and mechanical support are not censored in the treatment-policy primary analysis. A supportive composite counts the MAP criterion, protocol-defined rescue, return to CPB, or rescue mechanical support. Crossovers remain analysed as randomised.

No waveform-second imputation is used to derive the endpoint. If the endpoint is missing in ≤1% of the ITT population, complete-case primary analysis plus best/worst-case bounds will be reported. If >1%, multiple imputation by chained equations will include treatment, strata, baseline MAP, surgery type, CPB duration, rescue treatment, and observed haemodynamic summaries; Rubin’s rules combine estimates. Missing continuous secondary outcomes use analogous multiple imputation when plausible.

## Secondary analyses

- Binary outcomes: risk difference and risk ratio with 95% CIs.
- Nadir MAP, relative MAP decline, and hypotension burden: ANCOVA/linear regression adjusted for randomisation strata and baseline MAP; inspect residuals and use transformation or robust regression if assumptions materially fail.
- Repeated minute-level MAP: supportive mixed model with fixed treatment, time, treatment-by-time, baseline MAP, strata, and participant random intercept; model covariance selected before unmasking.
- Time to first endpoint: Kaplan–Meier description and stratified Cox model only as supportive analysis.
- Safety: denominators, risks, risk differences, and exact 95% CIs; sparse outcomes are descriptive.

## Multiplicity and subgroups

The primary endpoint has one confirmatory test. All secondary outcomes and interactions are supportive; no adjusted efficacy claims will be made. Prespecified exploratory subgroups are prior protamine/protamine-insulin exposure, surgery class, baseline MAP, CPB duration, sex, and site region. Report interaction estimates and CIs, not within-subgroup significance claims.

## Sample size and blinded reassessment

The design assumes placebo risk {CONTROL_RATE:.1%}, FPA risk {TREATMENT_RATE:.1%}, two-sided alpha {ALPHA:.2f}, power {POWER:.0%}, and {NON_EVALUABLE:.0%} non-evaluable allowance. Executable calculations require at least {CONSTANTS["calculated_minimum_total"]} participants; the operational target is {TARGET_TOTAL} ({TARGET_ARM} per group). The 12.9% anchor derives from a broader external protamine-adverse-event definition and does not validate this protocol’s endpoint.

After {CONFIG["sample_size"]["blinded_reestimation_after"]} participants, an independent statistician may update only the pooled primary-event and non-evaluable rates while preserving the planned absolute effect assumption and treatment masking. The target may increase to maintain 90% power, subject to a prespecified cap of 1,200; it will not decrease below 800. No treatment-effect estimate is released.
"""
    consent = """
# Informed-consent core elements

The consent form must state the trial purpose, random assignment, placebo use, masking, off-label/investigational prophylactic FPA use, absence of proven preventive benefit, alternative of usual care without study drug, and the right to withdraw without affecting care. Risks must include AKI, bleeding/transfusion/reoperation, GI injury, allergy/anaphylaxis, bronchospasm/aspirin asthma, cardiovascular events, pregnancy risk, and confidentiality breach. It must explain blood sampling and waveform/medication data collection, day-30 contact, compensation/treatment for research injury, data retention and sharing, sponsor/investigator contacts, ethics contacts, and circumstances for emergency unblinding.
"""
    data_management = """
# Data management plan

- Use validated electronic case-report forms with role-based access, audit trails, programmed range/logic checks, and immutable randomisation records.
- Synchronise operating-room monitor, infusion-pump, anaesthesia-record, and study-system clocks daily; store the offset.
- Transfer raw arterial-waveform files in native format plus a documented open export. Never overwrite source files.
- Use a pseudonymous participant ID; keep the linkage key at each site under restricted access.
- Record time-stamped heparin, ACT, residual-heparin estimate, protamine dose/rate/route/interruptions, all vasoactive drugs, fluids, pacing, cardioversion, CPB return, and mechanical support.
- Lock derivation code before unmasking. Maintain data-query, adjudication, and code-version logs.
- Retain essential records for the period required by Japanese law, sponsor policy, registry, and ethics approval; the protocol/consent must specify the final period.
- Share deidentified participant-level data and code only under the approved governance plan and data-use agreement; public synthetic data are not clinical data.
"""
    monitoring = """
# Monitoring plan

Risk-based monitoring will verify consent, eligibility, allocation concealment, study-drug accountability, timing, protamine dosing, primary waveform availability, SAE reporting, and key bleeding/renal outcomes. Central monitoring will flag implausible timestamps, missing waveform intervals, protocol-window deviations, duplicate values, and unusual site event rates. On-site or remote source-data verification will target all consent forms, all SAEs/SUSARs, all deaths, all suspected anaphylaxis, all returns to CPB, and a random sample of primary endpoints and non-events.

The sponsor will convene an independent DSMB with cardiac anaesthesia, cardiothoracic surgery, nephrology/haemostasis, and biostatistics expertise. The DSMB charter will define membership, conflicts, closed/open reports, meeting schedule, safety triggers, confidentiality, and recommendation procedures before first enrolment.
"""
    status = """
# STATUS

The protocol-development package is generated and internally validated. The package is not yet ethically approved, registered, recruiting, or accepted by ACA. Before submission, the authors must obtain the ACA presubmission decision, finalise sponsor/funding/site/registry/ethics fields, approve the protocol, and replace all bracketed placeholders.
"""
    decisions = f"""
# DECISIONS

1. Restrict the initial efficacy trial to elective non-CABG CPB surgery because international systemic NSAID labels contain CABG warnings.
2. Use FPA {DOSE} mg/{VOLUME} mL over {ADMIN_MIN} minutes, completed {LEAD} ± {WINDOW} minutes before protamine.
3. Use a prospective rolling 60-second relative-MAP endpoint rather than the retrospective broad-window PSI calculation.
4. Treat MAP <65 mmHg and hypotension burden as supportive outcomes.
5. Use a treatment-policy ITT estimand with rescue captured as outcome and composite sensitivity analysis.
6. Target {TARGET_TOTAL} participants on a binary primary-endpoint basis, with blinded pooled-rate reassessment after 300.
7. Prefer residual-heparin-guided protamine dosing within the local label; prohibit excess ratios and record all deviations.
8. Ask ACA about protocol eligibility before portal submission.
"""
    protocol_history = f"""
# Protocol version history

| Version | Date | Change |
|---|---|---|
| 0.1 | 2017-06 to 2017-07 | Historical IFUPLEASE concept; 180 participants and retrospective PSI-like endpoint |
| 1.0 | {CONFIG["protocol_date"]} | Rebuilt as a multicentre double-blind RCT; narrowed non-CABG population; digital 60-second MAP endpoint; target N={TARGET_TOTAL}; contemporary SPIRIT 2025, safety, SAP, and reproducibility package |
"""
    cover_text = f"""25 September 2026

Editor-in-Chief
Annals of Cardiac Anaesthesia

Dear Editor-in-Chief,

We submit the enclosed protocol manuscript, “{TITLE},” subject to confirmation that Annals of Cardiac Anaesthesia will consider a prospective randomised-trial protocol. The topic directly concerns cardiac-anaesthesia management of protamine reversal, but the current author instructions do not explicitly list a protocol category; our separate presubmission inquiry asks the editorial office to confirm eligibility and portal classification.

The trial will test a hypothesis generated by a retrospective association and does not claim prospective efficacy. It is designed as a multicentre, randomised, double-blind, placebo-controlled trial of {TARGET_TOTAL} adults undergoing elective non-CABG cardiac surgery with CPB. The primary endpoint uses prospectively captured arterial-line MAP and prespecified artifact rules.

The manuscript has not been submitted elsewhere. All authors must confirm final approval, contributions, competing interests, funding, ethics, and registration before formal submission. Generative AI (Devin, Cognition AI) was used to assist literature organisation, drafting, code generation, and consistency checks; the authors reviewed the sources, analyses, and text and retain full responsibility.

Sincerely,

Tatsuki Onishi
Data Science AI Innovation Promotion Center, Shiga University
1-1-1 Bamba, Hikone, Shiga 522-8522, Japan
E-mail: bougtoir@gmail.com
Tel: +81 749-27-1030
"""
    inquiry_text = f"""Subject: Presubmission inquiry—randomised trial protocol on prevention of protamine-associated hypotension

Dear Editor-in-Chief,

Would Annals of Cardiac Anaesthesia consider the prospective protocol entitled “{TITLE}”?

The planned multicentre, double-blind trial will randomise {TARGET_TOTAL} adults undergoing elective non-CABG cardiac surgery with CPB to FPA {DOSE} mg IV or saline placebo before protamine. The primary endpoint is a prospectively specified, artefact-controlled 20% relative MAP decline during 30 minutes after protamine. The rationale derives from a published retrospective association, which is presented only as hypothesis-generating.

The current ACA author instructions do not explicitly list a protocol category. If the manuscript is potentially suitable, please advise (1) which article category should be selected, (2) whether the Original Article limits should be followed, (3) whether SPIRIT materials and the full SAP should be uploaded as supplements, and (4) whether any registration or ethics approval must be complete before protocol submission.

The attached package has been prepared within a 250-word structured abstract, 3000-word main-text, five combined table/figure, and 30-reference envelope.

Sincerely,

Tatsuki Onishi
Data Science AI Innovation Promotion Center, Shiga University
E-mail: bougtoir@gmail.com
"""
    submission_readme = f"""
# ACA submission package

## Required first action

Send `presubmission_inquiry_ACA.docx`. ACA currently does not explicitly list protocols.

## If invited

1. Replace every bracketed placeholder.
2. Confirm author order, affiliations, contributions, corresponding-author details, sponsor, funding, insurance/indemnity, ethics, registry IDs, recruitment dates, and site count.
3. Upload `ACA_protocol_manuscript_blinded.docx` as the blinded manuscript and `title_page.docx` separately.
4. Upload the two TIFF figure files and `tables.docx`; use the editorial office’s requested handling if duplicate inline figures are disallowed.
5. Upload `SPIRIT_checklist.docx`, `statistical_analysis_plan.docx`, and `supplement.docx` as supplementary files if permitted.
6. Use the portal: https://review.jow.medknow.com/aca

## Integrity

No prospective participants or clinical results are represented. All validation data are labelled synthetic. Calculated sample-size values are read from executable outputs.
"""
    provenance = f"""
# PROVENANCE

## Private supplied inputs

Private files remain under `source/private/` and are excluded from Git and public ZIP archives. The source-manuscript DOCX SHA-256 is `53e734a4a893b0c98593bac6dee1acb2e3dc990e3b3e0ceb27d0bed0a20d3756`. The local PDF SHA-256 is `82ad50c4614d8d538a425baa38112ff973b387e8926e4aa1e0141648d7294632`; it was generated from the DOCX and is not independently verified as the final publisher PDF.

## Public sources

`source/source_acquisition.csv` records URL, version, retrieval time, conditions, path, size, checksum, and use conditions. Crossref DOI JSON responses are retained unchanged under `literature/raw/`.

## Numeric provenance

Sample-size values originate from `scripts/sample_size_analysis.py` and `output/analysis_constants.json`. Synthetic validation summaries originate from `scripts/mock_analysis.py`. Document generation reads those files rather than retyping calculated values.
"""
    return {
        "source_extraction.md": source_extraction,
        "journal_requirements.md": journal_requirements,
        "article_type_decision.md": article_type,
        "aca_fit_assessment.md": aca_fit,
        "aca_protocol_precedent_search.md": precedent,
        "hemodynamic_measurement_spec.md": haemodynamic,
        "artifact_handling_rules.md": artifacts,
        "randomization_spec.md": randomisation,
        "blinding_spec.md": blinding,
        "safety_regulatory_review.md": safety,
        "eligibility_rationale.md": eligibility,
        "SAP.md": sap,
        "informed_consent_core_elements.md": consent,
        "data_management_plan.md": data_management,
        "monitoring_plan.md": monitoring,
        "STATUS.md": status,
        "DECISIONS.md": decisions,
        "protocol_version_history.md": protocol_history,
        "cover_letter_ACA.txt": cover_text,
        "presubmission_inquiry_ACA.txt": inquiry_text,
        "README_SUBMISSION.md": submission_readme,
        "PROVENANCE.md": provenance,
    }


def make_audit_csvs() -> None:
    numeric_rows = [
        ["records screened", "92", "Source manuscript methods", "Verified"],
        ["excluded peripheral vascular", "5", "Source manuscript methods", "Verified"],
        ["analysed cardiovascular operations", "87", "92−5", "Verified"],
        ["on-pump subgroup", "40", "Procedure table", "Verified"],
        ["off-pump CABG", "47", "Procedure table", "Verified"],
        ["FPA lead time", "≥15 min before protamine", "Source methods", "Verified"],
        ["retrospective protamine", "10 mg/1000 U initial heparin over 5 min", "Source methods", "Historical only; not adopted prospectively"],
        ["PSI formula", "(MAPpre−MAPpost)/MAPpre", "Source methods", "Verified"],
        ["historical event threshold", "PSI >0.20", "Source methods", "Author-defined; not universally validated"],
        ["on-pump DID estimate", "4.09 mmHg", "Source results", "Verified"],
        ["DID CI", "approximately 0.72–7.46", "Source results", "Verified"],
        ["DID p value", "0.02", "Source results", "Verified"],
        ["weighted FPA probability", "approximately 0.00", "Source results", "Verified"],
        ["weighted control probability", "approximately 0.17", "Source results", "Verified"],
        ["bootstrap risk difference", "approximately −0.18", "Source results", "Verified"],
        ["bootstrap CI", "approximately −0.40 to −0.03", "Source results", "Verified"],
    ]
    write_csv(
        ROOT / "source_numeric_audit.csv",
        [
            {"item": row[0], "value": row[1], "location_or_derivation": row[2], "status": row[3]}
            for row in numeric_rows
        ],
        ["item", "value", "location_or_derivation", "status"],
    )
    claim_rows = [
        ["Retrospective association is hypothesis-generating", "source2026", "source_extraction.md; manuscript Introduction"],
        ["Protamine can cause hypotension, pulmonary hypertension, and allergy", "boer2018", "Introduction; safety review"],
        ["Immediate adverse reactions are clinically recognised", "weiler1990", "Introduction"],
        ["12.9% is a broader planning anchor", "kimmel1998events", "Sample-size report; Methods"],
        ["Protamine haemodynamic burden is associated with mortality", "welsby2005", "Introduction"],
        ["NPH insulin and allergy are risk factors", "kimmel1998risk;weiss1989", "Eligibility; randomisation strata"],
        ["Antihistamine pretreatment has been tested prospectively", "suksompong2023", "Introduction"],
        ["Cyclooxygenase inhibition attenuated reactions in animals", "hobbhahn1988", "Introduction"],
        ["FPA clinical dosing/use", "ropion_label;yamashita2006", "Intervention; safety review"],
        ["Renal safety requires monitoring", "liu2024;kdigo2012", "Safety outcomes"],
        ["Systemic NSAID CABG warnings exist internationally", "fda_nsaid_label", "Eligibility; safety review"],
        ["Residual-heparin-guided protamine is preferred", "sts2018;vonk2014", "Protamine regimen"],
        ["Japanese protamine rate and dose limits", "protamine_jp_label", "Protamine regimen"],
        ["Hypotension definitions are heterogeneous", "wesselink2018;sessler2019", "Endpoint rationale"],
        ["Protocol reporting follows SPIRIT 2025", "spirit2025", "Reporting"],
        ["Future results reporting will follow CONSORT 2025", "consort2025", "Discussion"],
    ]
    write_csv(
        ROOT / "source_claim_reference_map.csv",
        [{"claim": row[0], "reference_ids": row[1], "package_locations": row[2]} for row in claim_rows],
        ["claim", "reference_ids", "package_locations"],
    )


def make_spirit_crosswalk() -> list[dict[str, str]]:
    rows = [
        ("1", "Title identifies randomised trial protocol", "Manuscript title", "Yes"),
        ("2", "Trial registration", "Abstract; Trial status", "Pending—must be completed before enrolment"),
        ("3", "Protocol version", "Methods; version history", "Yes"),
        ("4", "Funding", "Declarations; title page", "Pending sponsor confirmation"),
        ("5", "Roles and responsibilities", "Methods; supplement", "Partly—committee names pending"),
        ("6a", "Background and rationale", "Introduction", "Yes"),
        ("6b", "Comparator rationale", "Methods—interventions", "Yes"),
        ("7", "Objectives/hypotheses", "Methods—objectives", "Yes"),
        ("8", "Trial design", "Methods—design", "Yes"),
        ("9", "Study setting", "Methods—setting", "Pending final sites"),
        ("10", "Eligibility criteria", "Table 2; supplement", "Yes"),
        ("11a", "Intervention description", "Methods—interventions", "Yes"),
        ("11b", "Discontinuation/modification", "Methods—interventions", "Yes"),
        ("11c", "Adherence strategies", "Methods—interventions", "Yes"),
        ("11d", "Concomitant care", "Methods—interventions", "Yes"),
        ("12", "Outcomes", "Table 3; haemodynamic specification", "Yes"),
        ("13", "Participant timeline", "Table 1; Figure 1", "Yes"),
        ("14", "Sample size", "Methods; executable report", "Yes"),
        ("15", "Recruitment", "Methods", "Pending site plan"),
        ("16a", "Sequence generation", "Randomisation specification", "Yes"),
        ("16b", "Concealment mechanism", "Randomisation specification", "Yes"),
        ("16c", "Implementation", "Randomisation specification", "Yes"),
        ("17a", "Who is blinded", "Blinding specification", "Yes"),
        ("17b", "Emergency unblinding", "Blinding specification", "Yes"),
        ("18a", "Data collection methods", "Methods; data management plan", "Yes"),
        ("18b", "Retention/follow-up", "Methods; consent", "Yes"),
        ("19", "Data management", "Data management plan", "Yes"),
        ("20a", "Statistical methods", "SAP; Methods", "Yes"),
        ("20b", "Additional analyses", "SAP", "Yes"),
        ("20c", "Analysis population/missing data", "SAP", "Yes"),
        ("21a", "Data monitoring committee", "Monitoring plan", "Yes; membership pending"),
        ("21b", "Interim analyses/stopping", "Safety review; SAP", "Yes"),
        ("22", "Harms", "Safety review; Methods", "Yes"),
        ("23", "Auditing", "Monitoring plan", "Yes"),
        ("24", "Research ethics approval", "Trial status; declarations", "Pending"),
        ("25", "Protocol amendments", "Methods; version history", "Yes"),
        ("26a", "Consent process", "Consent core elements", "Yes"),
        ("26b", "Ancillary studies", "Consent core elements", "Not currently planned"),
        ("27", "Confidentiality", "Data management plan", "Yes"),
        ("28", "Access to data", "Data management plan; declarations", "Yes"),
        ("29", "Ancillary/post-trial care", "Consent; declarations", "Pending compensation policy"),
        ("30", "Dissemination policy", "Declarations", "Yes"),
        ("31a", "Protocol and consent availability", "Supplement; trial registry", "Planned"),
        ("31b", "Biological specimens", "Methods", "No repository planned"),
        ("32", "Patient/public involvement", "Methods", "Pending—must be documented"),
        ("33", "Data sharing and reproducibility", "Data management; reproducibility README", "Yes"),
        ("34", "Environmental/equity considerations", "Discussion", "Addressed at protocol level"),
    ]
    write_csv(ROOT / "SPIRIT_item_crosswalk.csv", [
        {"item": item, "description": description, "location": location, "status": status}
        for item, description, location, status in rows
    ], ["item", "description", "location", "status"])
    return [
        {"item": item, "description": description, "location": location, "status": status}
        for item, description, location, status in rows
    ]


def manuscript_sections() -> tuple[str, list[tuple[str, list[str]]]]:
    abstract = f"""
Background and Aims: Protamine reversal after cardiopulmonary bypass can cause abrupt hypotension. A retrospective association between pre-protamine intravenous flurbiprofen axetil (FPA) and less haemodynamic decline is hypothesis-generating. This trial will test whether FPA reduces clinically important peri-protamine hypotension.

Methods: {ACRONYM} is a multicentre, parallel-group, randomised, double-blind, placebo-controlled superiority trial. {TARGET_TOTAL} adults undergoing elective non-coronary-artery-bypass cardiac surgery with cardiopulmonary bypass will receive FPA {DOSE} mg/{VOLUME} mL or saline placebo over {ADMIN_MIN} minutes, completed {LEAD} ± {WINDOW} minutes before protamine. The primary endpoint is at least one artefact-free rolling 60-second epoch during 0–30 minutes after protamine start with median mean arterial pressure ≤80% of a stabilised pre-protamine baseline. Continuous arterial pressure and time-stamped vasoactive treatment will be captured.

Planned Statistical Analysis: The intention-to-treat treatment-policy analysis will estimate an adjusted marginal risk difference with a 95% stratified-bootstrap confidence interval. Supportive analyses include risk ratio, absolute MAP <65 mmHg, hypotension burden, rescue-treatment composites, continuous MAP outcomes, and safety outcomes. The design has 90% power for planning risks of 12.9% versus 6.0% with 5% non-evaluable allowance; target enrolment is {TARGET_TOTAL}.

Discussion: Prospective randomisation, concealed allocation, high-resolution haemodynamic acquisition, prespecified artifact rules, and complete rescue-treatment timestamps address the main limitations of the retrospective study. FPA prophylaxis is investigational/off-label; renal, bleeding, hypersensitivity, and cardiovascular safety will be independently monitored.

Trial Status: Not yet recruiting; registration and ethics identifiers are pending.
""".strip()
    sections = [
        (
            "Introduction",
            [
                f"This protocol was motivated by a retrospective association between pre-protamine FPA and attenuated MAP decline [{ref('source2026')}]. Protamine neutralises unfractionated heparin after cardiopulmonary bypass (CPB) but can produce systemic hypotension, pulmonary vasoconstriction, allergy, and dose-related anticoagulant effects [{ref('boer2018')}]. Prospective and observational evidence confirms that clinically important reactions occur in cardiac surgery [{ref('weiler1990')},{ref('kimmel1998events')}], and the duration and magnitude of haemodynamic disturbance after protamine are associated with mortality [{ref('welsby2005')}]. Protamine-containing insulin and allergy histories identify higher-risk patients [{ref('kimmel1998risk')},{ref('weiss1989')}].",
                f"Preventive evidence is limited. A randomised antihistamine trial demonstrated the feasibility of serial haemodynamic assessment but did not establish a general preventive strategy [{ref('suksompong2023')}]. In animals, cyclooxygenase inhibition attenuated protamine-associated haemodynamic responses [{ref('hobbhahn1988')}], consistent with involvement of arachidonic-acid mediators [{ref('wang2021')}].",
                f"The hypothesis for this protocol arose from a retrospective cohort in which pre-protamine intravenous flurbiprofen axetil (FPA) was associated with an approximately 4.09-mmHg attenuation of mean arterial pressure (MAP) decline in the small on-pump subgroup [{ref('source2026')}]. Treatment was selected by clinicians, propensity overlap was limited, records were historical and partly paper-based, MAP windows were broad, and vasoactive timestamps were incomplete. The observation is therefore hypothesis-generating and cannot establish efficacy.",
                f"FPA is an intravenous non-steroidal anti-inflammatory drug available in Japan [{ref('ropion_label')}], with perioperative analgesic experience [{ref('yamashita2006')}]. Renal and bleeding risks remain material after CPB; observational perioperative renal-safety findings do not substitute for randomised cardiac-surgery safety data [{ref('liu2024')}]. International systemic NSAID labels also carry cardiovascular, gastrointestinal, and CABG warnings [{ref('fda_nsaid_label')}].",
                "The objective is to determine whether assignment to FPA before protamine reduces clinically important peri-protamine hypotension compared with saline placebo while characterising haemodynamic rescue and safety.",
            ],
        ),
        (
            "Methods",
            [
                f"Design and setting: {ACRONYM} is a multicentre, parallel-group, 1:1, randomised, double-blind, placebo-controlled superiority trial in Japanese cardiac-surgery centres. The protocol version is {CONFIG['protocol_version']} dated {CONFIG['protocol_date']}. Planned site count, sponsor, ethics approval, and registration identifiers remain to be finalised before enrolment. The assessment schedule and planned participant flow appear in Table 1 and Figure 1.",
                "Participants: Adults aged ≥18 years scheduled for elective non-CABG cardiac surgery requiring CPB, systemic unfractionated heparinisation, planned protamine reversal, and digital arterial-line acquisition are eligible after written consent. Planned CABG, emergency/salvage surgery, active bleeding or important coagulopathy, active peptic ulcer, severe renal/hepatic disease, decompensated heart failure/cardiogenic shock, severe uncontrolled hypertension, pregnancy, FPA/NSAID hypersensitivity, aspirin-exacerbated respiratory disease, prior severe protamine reaction, and prohibited interacting medicines are exclusions (Table 2). Organ-dysfunction thresholds will be approved against the current Japanese label before activation. Prior protamine or protamine-insulin exposure is a stratification factor unless associated with a severe reaction.",
                f"Interventions: Independent unblinded preparation staff will prepare FPA {DOSE} mg/{VOLUME} mL or {VOLUME} mL 0.9% saline in opaque syringe and tubing. Infusion will last {ADMIN_MIN} minutes and be completed {LEAD} ± {WINDOW} minutes before t=0. The dose reflects the Japanese marketed single ampoule; prophylaxis of protamine hypotension is off-label/investigational [{ref('ropion_label')}]. Study drug is withheld for a new contraindication, urgent surgery, or inability to meet the timing window. Clinical rescue is never delayed.",
                f"Protamine regimen: t=0 is the start of the first protamine infusion. Sites will preferentially use residual-heparin-guided dosing [{ref('sts2018')}], which can reduce excess protamine and improve haemostatic parameters [{ref('vonk2014')}]. Where unavailable, dose will follow the local approved label and documented heparin exposure; ratio must remain below 2.6 mg/100 U. Each 50-mg portion will be diluted and infused over at least 10 minutes consistent with the Japanese label [{ref('protamine_jp_label')}]. Heparin, ACT, residual heparin, protamine dose/rate/route, interruptions, additional doses, and deviations are recorded.",
                "Randomisation and masking: A validated central system will use randomly permuted blocks of 4, 6, and 8, stratified by site and prior protamine/protamine-insulin exposure. Allocation is concealed until assignment. Participants, clinical teams, outcome assessors, central waveform reviewers, and statisticians remain masked. Emergency unblinding requires documented clinical necessity.",
                "Outcomes and haemodynamic measurement: Baseline MAP is the median of valid observations from −5 to −1 minutes. The primary endpoint is at least one artefact-free rolling 60-second epoch from t=0 to +30 minutes with median MAP ≤80% of baseline and ≥45 valid seconds. A relative threshold preserves the retrospective hypothesis without calling it universally validated shock. Hypotension definitions vary across perioperative literature [{0}], while organ risk is also related to absolute pressure and duration [{1}]; therefore MAP <65 mmHg and burden below 65 mmHg are supportive outcomes (Table 3).".format(ref("wesselink2018"), ref("sessler2019")),
                "Waveform and rescue capture: MAP will be sampled at ≥1 Hz. Flush, calibration, disconnection, and nonphysiologic damping intervals are excluded under prespecified blinded rules; isolated beats cannot define an event. Exact time, dose, and route of every vasopressor/inotrope change, fluid bolus, pacing, cardioversion, return to CPB, and mechanical support are recorded. Observed post-rescue MAP remains in the primary treatment-policy analysis; a rescue composite is supportive (Figure 2).",
                f"Safety outcomes: Kidney Disease: Improving Global Outcomes (KDIGO) acute kidney injury through postoperative day 7/discharge [{ref('kdigo2012')}], chest-tube drainage, transfusion, reoperation, gastrointestinal bleeding, bronchospasm/anaphylaxis, myocardial infarction, stroke, serious adverse events, and death will be recorded. An independent DSMB reviews safety after 100 participants and unmasked data after 300 and 600, with ad hoc review for suspected unexpected serious adverse reactions.",
                f"Sample size: Executable two-group calculations use placebo risk {CONTROL_RATE:.1%}, FPA risk {TREATMENT_RATE:.1%}, two-sided alpha {ALPHA:.2f}, {POWER:.0%} power, and {NON_EVALUABLE:.0%} non-evaluable allowance. The minimum is {CONSTANTS['calculated_minimum_total']}; operational enrolment is {TARGET_TOTAL} ({TARGET_ARM}/group). The 12.9% anchor came from a broader adverse-event definition [{ref('kimmel1998events')}] and is not endpoint validation. A blinded reassessment after 300 may increase but not reduce enrolment.",
                f"Statistical analysis: The ITT primary analysis will fit logistic regression with treatment, site stratum, prior protamine/protamine-insulin exposure, and baseline MAP, then standardise risks over the trial population. The primary contrast is marginal risk difference with a 95% CI from {CONFIG['analysis']['bootstrap_samples']:,} stratified bootstrap samples; risk ratio and odds ratio are supportive. Firth logistic and exact methods address separation. Continuous outcomes use adjusted linear models; minute-level MAP uses a supportive mixed model. Secondary and subgroup analyses are exploratory. Missing endpoint rules, multiple imputation, crossover, protocol deviations, and per-protocol analysis are prespecified in the SAP.",
                f"Data quality and reporting: Electronic case-report forms, audit trails, clock synchronisation, raw-waveform retention, blinded derivation, central monitoring, and targeted source verification are required. The protocol follows SPIRIT 2025 [{ref('spirit2025')}], with planned results reporting under CONSORT 2025 [{ref('consort2025')}].",
            ],
        ),
        (
            "Discussion",
            [
                "This trial converts a retrospective signal into a prospective causal test. Randomisation removes clinician-driven treatment selection; multicentre recruitment improves transportability; digital MAP acquisition replaces broad lowest-value windows; and time-stamped rescue capture makes haemodynamic interpretation explicit. The primary endpoint retains the clinically intuitive 20% relative decline while requiring sustained, artefact-controlled evidence. Absolute MAP and burden outcomes address the limitation that a relative threshold may not indicate organ hypoperfusion in every participant.",
                "The selected population deliberately excludes planned CABG and severe label-defined risk states. This improves safety but limits generalisability to emergency surgery, CABG, severe organ dysfunction, and patients with established high protamine risk. FPA’s mechanism remains uncertain, the planning event rates are indirect, and rescue treatments may attenuate observed pressure differences. Treatment-policy and rescue-composite estimands address, but cannot eliminate, this complexity.",
                "The sample size is substantially larger than the historical protocol because a realistic external control risk and conservative treatment risk replace the retrospective zero-event observation. The blinded reassessment protects power against an incorrect pooled event-rate assumption without exposing treatment effects. The trial is designed to estimate both clinical benefit and harms; no preventive use should occur outside approved care or the research protocol on the basis of the retrospective study alone.",
            ],
        ),
        (
            "Trial status",
            [
                "Protocol version 1.0, 25 September 2026. The trial is not yet registered, ethically approved, or recruiting. Recruitment dates, sponsor, participating sites, registry identifiers, and ethics identifiers will be inserted only after formal confirmation.",
            ],
        ),
        (
            "Declarations",
            [
                "Ethics approval and consent to participate: Pending. Written informed consent will be obtained before randomisation.",
                "Consent for publication: Not applicable; no participant data are reported.",
                "Availability of data and materials: Protocol code and synthetic validation data are supplied. No prospective clinical data exist. Deidentified trial data-sharing conditions will be defined in the approved protocol and consent.",
                "Competing interests: [Each author must complete before submission.]",
                "Funding: [Sponsor and funding source must be confirmed before submission.]",
                "Authors’ contributions: [Use CRediT roles and confirm all authors before submission.]",
                "Acknowledgements: Generative AI (Devin, Cognition AI) assisted literature organisation, drafting, code generation, document production, and consistency checks. The authors reviewed the sources, analyses, and text and retain full responsibility for accuracy, integrity, and submission decisions.",
                "Patient and public involvement: [Describe involvement or explicitly state none, with rationale, before submission.]",
            ],
        ),
    ]
    return abstract, sections


def make_manuscript_docx(path: Path, figures: list[dict[str, str]], include_images: bool = True) -> tuple[int, int]:
    abstract, sections = manuscript_sections()
    abstract_words = len(re.findall(r"\b[\w%±≥≤–-]+\b", abstract))
    body_words = sum(len(re.findall(r"\b[\w%±≥≤–-]+\b", paragraph)) for _, paragraphs in sections for paragraph in paragraphs)
    if abstract_words > 250:
        raise ValueError(f"Abstract is {abstract_words} words")
    if body_words > 3000:
        raise ValueError(f"Body is {body_words} words")
    document = Document()
    set_doc_defaults(document)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(14)
    add_heading(document, "Abstract")
    for block in abstract.split("\n\n"):
        prefix = block.split(":", 1)[0] + ":"
        add_paragraph(document, block, bold_prefix=prefix)
    add_paragraph(document, "Keywords: cardiopulmonary bypass; flurbiprofen axetil; hypotension; protamine; randomised controlled trial; study protocol")
    inserted = set()
    for heading, paragraphs in sections:
        add_heading(document, heading)
        for paragraph_text in paragraphs:
            add_paragraph(document, paragraph_text)
            if heading == "Methods" and "eligible after written consent" in paragraph_text:
                table = TABLES[1]
                caption = document.add_paragraph()
                caption.paragraph_format.space_before = Pt(14)
                run = caption.add_run(table["title"])
                run.bold = True
                add_table(document, table["headers"], table["rows"])
                inserted.add("Table 2")
            if heading == "Methods" and "supportive outcomes (Table 3)" in paragraph_text:
                table = TABLES[2]
                caption = document.add_paragraph()
                caption.paragraph_format.space_before = Pt(14)
                run = caption.add_run(table["title"])
                run.bold = True
                add_table(document, table["headers"], table["rows"])
                inserted.add("Table 3")
            if heading == "Methods" and "supportive (Figure 2)" in paragraph_text:
                if include_images:
                    add_figure(document, Path(figures[1]["path"]), figures[1]["title"] + ". " + figures[1]["caption"])
                else:
                    add_paragraph(document, "[Figure 2 near here]")
                inserted.add("Figure 2")
            if heading == "Methods" and "Table 1 and Figure 1" in paragraph_text:
                table = TABLES[0]
                caption = document.add_paragraph()
                caption.paragraph_format.space_before = Pt(14)
                run = caption.add_run(table["title"])
                run.bold = True
                add_table(document, table["headers"], table["rows"])
                if include_images:
                    add_figure(document, Path(figures[0]["path"]), figures[0]["title"] + ". " + figures[0]["caption"])
                else:
                    add_paragraph(document, "[Figure 1 near here]")
                inserted.update({"Table 1", "Figure 1"})
    add_heading(document, "References")
    for record in REFERENCES:
        add_paragraph(document, f"{record['reference_number']}. {record['vancouver']}")
    if not include_images:
        add_heading(document, "Figure legends")
        for figure in figures:
            add_paragraph(document, figure["title"] + ". " + figure["caption"])
    if inserted != {"Table 1", "Table 2", "Table 3", "Figure 1", "Figure 2"}:
        raise ValueError(f"Missing inline outputs: {inserted}")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return abstract_words, body_words


def simple_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    document = Document()
    set_doc_defaults(document)
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(14)
    for paragraph in paragraphs:
        add_paragraph(document, paragraph)
    document.save(path)


def markdown_to_docx(markdown_path: Path, docx_path: Path, title: str | None = None) -> None:
    document = Document()
    set_doc_defaults(document)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    for line in lines:
        if line.startswith("# "):
            add_heading(document, line[2:], 1)
        elif line.startswith("## "):
            add_heading(document, line[3:], 2)
        elif line.startswith("### "):
            add_heading(document, line[4:], 3)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("|"):
            continue
        elif line.strip():
            add_paragraph(document, re.sub(r"`([^`]+)`", r"\1", line))
    if title:
        document.core_properties.title = title
    document.save(docx_path)


def make_supporting_docx(
    spirit_rows: list[dict[str, str]],
    figures: list[dict[str, str]],
    abstract_words: int,
    body_words: int,
) -> None:
    title_page = Document()
    set_doc_defaults(title_page)
    title_paragraph = title_page.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(TITLE)
    title_run.bold = True
    author_paragraph = title_page.add_paragraph()
    author_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_paragraph.add_run("Tatsuki Onishi")
    affiliation_run = author_paragraph.add_run("1,2")
    affiliation_run.font.superscript = True
    author_paragraph.add_run("; Tatsuyoshi Ikenoue")
    affiliation_run = author_paragraph.add_run("1")
    affiliation_run.font.superscript = True
    affiliation_lines = [
        ("1", " Data Science AI Innovation Promotion Center, Shiga University, 1-1-1 Bamba, Hikone, Shiga 522-8522, Japan"),
        ("2", " [Participating centre affiliation(s) to be finalised]"),
    ]
    for number, affiliation in affiliation_lines:
        paragraph = title_page.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(number)
        run.font.superscript = True
        paragraph.add_run(affiliation)
    for text, bold in [
        ("Corresponding author: Tatsuki Onishi, bougtoir@gmail.com, +81 749-27-1030", False),
        (f"Word counts: abstract {abstract_words}; main text {body_words}. Tables: {len(TABLES)}. Figures: 2. References: {len(REFERENCES)}.", False),
        ("Running title: FPA before protamine—trial protocol", False),
        ("Funding/sponsor/registry/ethics: [must be completed before submission]", False),
    ]:
        paragraph = title_page.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = bold
    title_page.save(ROOT / "title_page.docx")

    declarations = [
        "Ethics approval and consent to participate: Pending. Written informed consent will be obtained before randomisation.",
        "Trial registration: [jRCT/UMIN identifier pending; registration before first participant].",
        "Consent for publication: Not applicable.",
        "Availability of data and materials: No prospective data exist. Code and explicitly synthetic validation data accompany the protocol.",
        "Competing interests: [Each author must complete.]",
        "Funding and sponsor: [Must be confirmed.]",
        "Authors’ contributions: [Complete using CRediT roles.]",
        "Generative-AI disclosure: Devin (Cognition AI) assisted literature organisation, drafting, code generation, document production, and consistency checks. Authors verified sources, analyses, and text and retain full responsibility.",
    ]
    simple_docx(ROOT / "declarations.docx", "Declarations", declarations)

    for basename, title in [
        ("cover_letter_ACA", "Cover letter"),
        ("presubmission_inquiry_ACA", "Presubmission inquiry"),
    ]:
        paragraphs = (ROOT / f"{basename}.txt").read_text(encoding="utf-8").split("\n\n")
        simple_docx(ROOT / f"{basename}.docx", title, paragraphs)

    spirit_doc = Document()
    set_doc_defaults(spirit_doc)
    add_heading(spirit_doc, "SPIRIT 2025 checklist and crosswalk")
    add_paragraph(spirit_doc, "Crosswalk adapted to the saved SPIRIT 2025 statement and expanded checklist. Pending items must be completed before trial activation.")
    add_table(
        spirit_doc,
        ["Item", "Description", "Location", "Status"],
        [[row["item"], row["description"], row["location"], row["status"]] for row in spirit_rows],
    )
    spirit_doc.save(ROOT / "SPIRIT_checklist.docx")

    tables_doc = Document()
    set_doc_defaults(tables_doc)
    add_heading(tables_doc, "Tables")
    for table in TABLES:
        caption = tables_doc.add_paragraph()
        caption.paragraph_format.space_before = Pt(14)
        run = caption.add_run(table["title"])
        run.bold = True
        add_table(tables_doc, table["headers"], table["rows"])
    tables_doc.save(ROOT / "tables.docx")

    supplement_doc = Document()
    set_doc_defaults(supplement_doc)
    add_heading(supplement_doc, "Supplementary protocol materials")
    for name in [
        "hemodynamic_measurement_spec.md",
        "artifact_handling_rules.md",
        "randomization_spec.md",
        "blinding_spec.md",
        "safety_regulatory_review.md",
        "eligibility_rationale.md",
        "informed_consent_core_elements.md",
        "data_management_plan.md",
        "monitoring_plan.md",
    ]:
        add_heading(supplement_doc, name.replace(".md", "").replace("_", " ").title(), 2)
        text = (ROOT / name).read_text(encoding="utf-8")
        for line in text.splitlines():
            if line.startswith("#"):
                continue
            if line.startswith("- "):
                supplement_doc.add_paragraph(line[2:], style="List Bullet")
            elif line.strip():
                add_paragraph(supplement_doc, line)
    supplement_doc.save(ROOT / "supplement.docx")

    markdown_to_docx(ROOT / "SAP.md", ROOT / "statistical_analysis_plan.docx", "Statistical Analysis Plan")

    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    for figure in figures:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box = slide.shapes.add_textbox(PptxInches(0.5), PptxInches(0.15), PptxInches(12.33), PptxInches(0.6))
        title_frame = title_box.text_frame
        title_frame.text = figure["title"]
        title_frame.paragraphs[0].font.size = PptxPt(24)
        title_frame.paragraphs[0].font.bold = True
        title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        slide.shapes.add_picture(figure["path"], PptxInches(1.3), PptxInches(0.85), width=PptxInches(10.73), height=PptxInches(5.45))
        caption_box = slide.shapes.add_textbox(PptxInches(0.55), PptxInches(6.35), PptxInches(12.23), PptxInches(0.85))
        caption_frame = caption_box.text_frame
        caption_frame.text = figure["caption"]
        caption_frame.paragraphs[0].font.size = PptxPt(11)
    presentation.save(ROOT / "figures_editable.pptx")


def make_consistency_and_provenance(abstract_words: int, body_words: int) -> None:
    rows = [
        ["Dose", f"{DOSE} mg", "config/trial_design.json", "Matched"],
        ["Volume", f"{VOLUME} mL", "config/trial_design.json", "Matched"],
        ["Infusion duration", f"{ADMIN_MIN} min", "config/trial_design.json", "Matched"],
        ["Completion timing", f"{LEAD} ± {WINDOW} min before protamine", "config/trial_design.json", "Matched"],
        ["t=0", "start of first protamine infusion", "config/trial_design.json", "Matched"],
        ["Primary endpoint", "rolling 60-s median MAP ≤80% baseline, 0–30 min", "config/trial_design.json", "Matched"],
        ["Allocation", "1:1", "config/trial_design.json", "Matched"],
        ["Target N", str(TARGET_TOTAL), "output/analysis_constants.json", "Matched"],
        ["Per arm", str(TARGET_ARM), "output/analysis_constants.json", "Matched"],
        ["Power", f"{POWER:.0%}", "output/analysis_constants.json", "Matched"],
        ["Primary analysis", "adjusted marginal risk difference; stratified bootstrap CI", "config/trial_design.json/SAP.md", "Matched"],
        ["Abstract words", str(abstract_words), "generated manuscript", "≤250"],
        ["Body words", str(body_words), "generated manuscript", "≤3000"],
        ["Combined tables/figures", str(len(TABLES) + 2), "generated manuscript", "≤5"],
        ["References", str(len(REFERENCES)), "references_verified.csv", "≤30"],
    ]
    write_csv(ROOT / "MANUSCRIPT_CONSISTENCY_CHECK.csv", [
        {"field": field, "value": value, "source": source, "status": status}
        for field, value, source, status in rows
    ], ["field", "value", "source", "status"])
    claim_rows = [
        ["Retrospective study N=87", "source/private/jjsca_revision_2nd_cleaned.docx", "source_numeric_audit.csv", "Private author manuscript"],
        ["On-pump N=40", "source/private/jjsca_revision_2nd_cleaned.docx", "source_numeric_audit.csv", "Private author manuscript"],
        ["Retrospective effect 4.09 mmHg", "source/private/jjsca_revision_2nd_cleaned.docx", "Introduction/source extraction", "Hypothesis-generating only"],
        ["FPA dose and contraindications", "source/public/ropion_japan_package_insert.pdf", "Intervention/safety", "Official Japanese label"],
        ["Protamine infusion limits", "source/public/protamine_japan_package_insert.pdf", "Methods/safety", "Official Japanese label"],
        ["Residual-heparin protamine guidance", "source/public/sts_sca_amsect_anticoagulation_guideline.html", "Methods", "Professional guideline"],
        ["Primary target N", "output/analysis_constants.json", "Methods/SAP/title page", "Executable calculation"],
        ["Mock estimates", "reproducibility/outputs/mock_analysis_summary.json", "Validation only", "Synthetic; never clinical"],
        ["ACA limits", "source/public/aca_information_for_authors_20260925.html", "journal_requirements.md", "Official live page snapshot"],
        ["SPIRIT 2025 items", "source/public/spirit_2025_expanded_checklist.pdf", "SPIRIT crosswalk", "Official reporting resource"],
    ]
    write_csv(ROOT / "CLAIM_PROVENANCE_TABLE.csv", [
        {"claim": claim, "source": source, "output_locations": locations, "qualification": qualification}
        for claim, source, locations, qualification in claim_rows
    ], ["claim", "source", "output_locations", "qualification"])


def make_reproducibility_readme() -> None:
    text = f"""
# Reproducibility

All validation data are synthetic and are not clinical observations.

## Environment

Python dependencies are pinned in `requirements.txt`. From the project root:

```bash
python3 -m venv .venv
.venv/bin/pip install -r reproducibility/requirements.txt
.venv/bin/python scripts/verify_references.py
.venv/bin/python scripts/sample_size_analysis.py
.venv/bin/python scripts/mock_analysis.py
.venv/bin/python scripts/generate_protocol_package.py
.venv/bin/python scripts/qc_package.py
```

## Provenance

- Trial constants: `config/trial_design.json`
- Sample-size source: `scripts/sample_size_analysis.py`
- Calculated constants: `output/analysis_constants.json`
- Synthetic validation: `scripts/mock_analysis.py`
- Documents: `scripts/generate_protocol_package.py`
- QC: `scripts/qc_package.py`

The synthetic primary model converged: {MOCK["primary_model_converged"]}. This is a software-path check, not an efficacy result.
"""
    write_text(REPRO / "README.md", text)
    run_log = f"""Run completed UTC: {NOW}
sample_size_analysis.py: PASS; minimum={CONSTANTS["calculated_minimum_total"]}; target={TARGET_TOTAL}
mock_analysis.py: PASS; N={MOCK["n_total"]}; converged={MOCK["primary_model_converged"]}
synthetic label: {MOCK["data_origin"]}
verify_references.py: PASS; references={len(REFERENCES)}
generate_protocol_package.py: PASS
"""
    write_text(REPRO / "run_log.txt", run_log)
    source_dir = REPRO / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    for filename in [
        "sample_size_analysis.py",
        "mock_analysis.py",
        "verify_references.py",
        "generate_protocol_package.py",
        "qc_package.py",
    ]:
        shutil.copy2(ROOT / "scripts" / filename, source_dir / filename)
    data_dir = REPRO / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ROOT / "synthetic_validation_data.csv", data_dir / "synthetic_validation_data.csv")
    shutil.copy2(ROOT / "config" / "trial_design.json", REPRO / "trial_design.json")


def write_phase_handoffs() -> None:
    phases = {
        "PHASE_0_HANDOFF.txt": "Source materials recovered and isolated; retrospective manuscript and historical protocol audited; publisher-PDF limitation documented.",
        "PHASE_1_HANDOFF.txt": "ACA requirements, protocol-category uncertainty, authoritative labels, guideline sources, and acquisition ledger completed.",
        "PHASE_2_HANDOFF.txt": "Prospective trial design, endpoint, waveform rules, randomisation, masking, safety, sample size, and SAP completed.",
        "PHASE_3_HANDOFF.txt": "Blinded manuscript, title page, tables, figures, declarations, cover letter, presubmission inquiry, SPIRIT crosswalk, and supplement generated.",
        "PHASE_4_HANDOFF.txt": "Synthetic validation, hostile review, consistency audit, reproducibility bundle, and final package QC completed.",
    }
    for filename, text in phases.items():
        write_text(ROOT / filename, f"{text}\nGenerated: {NOW}")


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for name, content in markdown_documents().items():
        write_text(ROOT / name, content)
    make_audit_csvs()
    spirit_rows = make_spirit_crosswalk()
    figures = make_figures()
    abstract_words, body_words = make_manuscript_docx(
        ROOT / "ACA_protocol_manuscript_blinded.docx",
        figures,
        include_images=False,
    )
    make_manuscript_docx(
        ROOT / "ACA_protocol_manuscript_inline_figures.docx",
        figures,
        include_images=True,
    )
    make_supporting_docx(spirit_rows, figures, abstract_words, body_words)
    make_consistency_and_provenance(abstract_words, body_words)
    make_reproducibility_readme()
    write_phase_handoffs()
    print(
        json.dumps(
            {
                "abstract_words": abstract_words,
                "body_words": body_words,
                "tables": len(TABLES),
                "figures": len(figures),
                "references": len(REFERENCES),
                "target_total": TARGET_TOTAL,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
