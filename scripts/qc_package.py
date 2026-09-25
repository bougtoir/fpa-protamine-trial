#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
REPRO = ROOT / "reproducibility"
OUTPUT = ROOT / "output"
NOW = datetime.now(timezone.utc).replace(microsecond=0).isoformat()

REQUIRED_DOCX = [
    "ACA_protocol_manuscript_blinded.docx",
    "ACA_protocol_manuscript_inline_figures.docx",
    "title_page.docx",
    "declarations.docx",
    "cover_letter_ACA.docx",
    "presubmission_inquiry_ACA.docx",
    "SPIRIT_checklist.docx",
    "tables.docx",
    "supplement.docx",
    "statistical_analysis_plan.docx",
]

REPRO_FILES = [
    "README.md",
    "requirements.txt",
    "run_log.txt",
    "trial_design.json",
    "data/synthetic_validation_data.csv",
    "data/synthetic_map_data.csv",
    "outputs/mock_analysis_summary.json",
    "outputs/mock_primary_analysis.csv",
    "outputs/edge_case_checks.csv",
    "source/sample_size_analysis.py",
    "source/mock_analysis.py",
    "source/verify_references.py",
    "source/generate_protocol_package.py",
    "source/qc_package.py",
]

SUBMISSION_FILES = [
    "README_SUBMISSION.md",
    "ACA_protocol_manuscript_blinded.docx",
    "ACA_protocol_manuscript_inline_figures.docx",
    "title_page.docx",
    "declarations.docx",
    "cover_letter_ACA.docx",
    "cover_letter_ACA.txt",
    "presubmission_inquiry_ACA.docx",
    "presubmission_inquiry_ACA.txt",
    "SPIRIT_checklist.docx",
    "SPIRIT_item_crosswalk.csv",
    "tables.docx",
    "figures_editable.pptx",
    "output/figures/Figure_1_trial_timeline.png",
    "output/figures/Figure_1_trial_timeline.tiff",
    "output/figures/Figure_2_hemodynamic_timeline.png",
    "output/figures/Figure_2_hemodynamic_timeline.tiff",
    "statistical_analysis_plan.docx",
    "supplement.docx",
    "references_verified.csv",
    "MANUSCRIPT_CONSISTENCY_CHECK.csv",
    "CLAIM_PROVENANCE_TABLE.csv",
    "HOSTILE_REVIEW.md",
    "FINAL_QC_REPORT.md",
    "reproducibility_bundle.zip",
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def docx_text(path: Path) -> str:
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(cell.text for cell in row.cells)
    return "\n".join(paragraphs)


def check_docx(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "[Content_Types].xml" not in names or "word/document.xml" not in names:
            raise ValueError(f"Invalid DOCX: {path}")
    Document(path)


def first_citation_order(text: str) -> list[int]:
    order: list[int] = []
    for group in re.findall(r"\[((?:\d+,?)+)\]", text):
        for value in group.split(","):
            number = int(value)
            if number not in order:
                order.append(number)
    return order


def make_zip(zip_path: Path, entries: list[tuple[Path, str]]) -> None:
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for source, arcname in entries:
            if not source.exists():
                raise FileNotFoundError(source)
            archive.write(source, arcname)
    with zipfile.ZipFile(zip_path) as archive:
        bad = archive.testzip()
        if bad:
            raise ValueError(f"Corrupt entry {bad} in {zip_path}")


def main() -> None:
    results: list[tuple[str, str, str]] = []
    for filename in REQUIRED_DOCX:
        path = ROOT / filename
        check_docx(path)
        results.append(("DOCX opens", filename, "PASS"))

    blinded = ROOT / "ACA_protocol_manuscript_blinded.docx"
    inline = ROOT / "ACA_protocol_manuscript_inline_figures.docx"
    with zipfile.ZipFile(blinded) as archive:
        embedded = [name for name in archive.namelist() if name.startswith("word/media/")]
    if embedded:
        raise ValueError("Blinded submission manuscript unexpectedly embeds figures")
    results.append(("Submission figure separation", "No word/media in blinded manuscript", "PASS"))
    with zipfile.ZipFile(inline) as archive:
        embedded_inline = [name for name in archive.namelist() if name.startswith("word/media/")]
    if len(embedded_inline) != 2:
        raise ValueError("Inline manuscript does not contain exactly two figures")
    results.append(("Inline review manuscript", "Two embedded figures", "PASS"))

    text = docx_text(blinded)
    order = first_citation_order(text)
    expected = list(range(1, 23))
    if order != expected:
        raise ValueError(f"Citation order is {order}, expected {expected}")
    results.append(("Vancouver order", "References 1–22 first cited sequentially", "PASS"))

    for label in ["Table 1", "Table 2", "Table 3", "Figure 1", "Figure 2"]:
        if label not in text:
            raise ValueError(f"Missing citation or legend for {label}")
    results.append(("Table/figure cross-references", "Tables 1–3 and Figures 1–2 cited", "PASS"))
    table_first = [text.index(f"Table {number}") for number in (1, 2, 3)]
    figure_first = [text.index(f"Figure {number}") for number in (1, 2)]
    if table_first != sorted(table_first) or figure_first != sorted(figure_first):
        raise ValueError("Tables or figures are not first cited in numeric order")
    results.append(("Table/figure order", "First citations are sequential", "PASS"))

    consistency = {row["field"]: row for row in csv.DictReader((ROOT / "MANUSCRIPT_CONSISTENCY_CHECK.csv").open())}
    if int(consistency["Abstract words"]["value"]) > 250:
        raise ValueError("Abstract exceeds 250 words")
    if int(consistency["Body words"]["value"]) > 3000:
        raise ValueError("Body exceeds 3000 words")
    if int(consistency["Combined tables/figures"]["value"]) > 5:
        raise ValueError("More than five combined tables/figures")
    if int(consistency["References"]["value"]) > 30:
        raise ValueError("More than 30 references")
    results.append(("ACA numeric limits", "Abstract, body, displays, references", "PASS"))

    references = list(csv.DictReader((ROOT / "references_verified.csv").open()))
    if len(references) != 22 or any(not row["verification_status"] for row in references):
        raise ValueError("Reference verification incomplete")
    results.append(("Reference audit", "22 references with verification status", "PASS"))

    acquisition = list(csv.DictReader((ROOT / "source" / "source_acquisition.csv").open()))
    for row in acquisition:
        path = ROOT / row["local_path"]
        if not path.exists() or path.stat().st_size != int(row["bytes"]) or sha256(path) != row["sha256"]:
            raise ValueError(f"Source ledger mismatch: {row['local_path']}")
    results.append(("Source acquisition ledger", f"{len(acquisition)} persisted sources verified", "PASS"))

    synthetic = ROOT / "synthetic_validation_data.csv"
    header = synthetic.read_text(encoding="utf-8").splitlines()[0]
    sample = synthetic.read_text(encoding="utf-8").splitlines()[1]
    if "data_origin" not in header or "SYNTHETIC VALIDATION DATA" not in sample:
        raise ValueError("Synthetic participant data are not clearly labelled")
    map_header = (REPRO / "data" / "synthetic_map_data.csv").read_text(encoding="utf-8").splitlines()[0]
    if "data_origin" not in map_header:
        raise ValueError("Synthetic MAP data are not clearly labelled")
    results.append(("Fabrication safeguard", "Synthetic files explicitly labelled; manuscript reports no results", "PASS"))

    forbidden = [
        "we demonstrate that fpa prevents",
        "the trial showed",
        "participants were enrolled",
        "prospective efficacy was",
    ]
    lowered = text.lower()
    if any(phrase in lowered for phrase in forbidden):
        raise ValueError("Prospective efficacy/enrolment language detected")
    results.append(("Prospective language", "No completed-trial efficacy or enrolment claim detected", "PASS"))

    mock = json.loads((REPRO / "outputs" / "mock_analysis_summary.json").read_text())
    if mock["data_origin"] != "SYNTHETIC VALIDATION DATA—NOT CLINICAL DATA" or not mock["primary_model_converged"]:
        raise ValueError("Mock analysis validation failed")
    edge = pd_read_csv(REPRO / "outputs" / "edge_case_checks.csv")
    if not edge or any(row.get("status") != "converged" for row in edge):
        raise ValueError("Sparse-event edge cases failed")
    results.append(("Executable mock analysis", "Primary model converged and sparse-event checks passed", "PASS"))
    standalone = ROOT / "STANDALONE_REPRODUCTION_CHECK.md"
    if standalone.exists() and "Result: PASS" in standalone.read_text(encoding="utf-8"):
        results.append(("Standalone clean rebuild", "Extracted archive passed make all; numeric outputs and manuscript text matched", "PASS"))

    for filename in REPRO_FILES:
        if not (REPRO / filename).exists():
            raise FileNotFoundError(REPRO / filename)

    repro_entries = [
        (ROOT / "README.md", "README.md"),
        (ROOT / "Makefile", "Makefile"),
        (REPRO / "README.md", "reproducibility/README.md"),
        (REPRO / "requirements.txt", "reproducibility/requirements.txt"),
        (REPRO / "run_log.txt", "reproducibility/run_log.txt"),
        (REPRO / "trial_design.json", "reproducibility/trial_design.json"),
        (REPRO / "data" / "synthetic_validation_data.csv", "reproducibility/data/synthetic_validation_data.csv"),
        (REPRO / "data" / "synthetic_map_data.csv", "reproducibility/data/synthetic_map_data.csv"),
        (REPRO / "outputs" / "mock_analysis_summary.json", "reproducibility/outputs/mock_analysis_summary.json"),
        (REPRO / "outputs" / "mock_primary_analysis.csv", "reproducibility/outputs/mock_primary_analysis.csv"),
        (REPRO / "outputs" / "edge_case_checks.csv", "reproducibility/outputs/edge_case_checks.csv"),
    ]
    for filename in [
        "sample_size_analysis.py",
        "mock_analysis.py",
        "verify_references.py",
        "generate_protocol_package.py",
        "qc_package.py",
        "browser_capture.py",
    ]:
        repro_entries.append((ROOT / "scripts" / filename, f"scripts/{filename}"))
    for filename in [
        "sample_size_results.csv",
        "sample_size_report.md",
        "sample_size_figure.png",
        "output/analysis_constants.json",
        "output/references_vancouver.json",
        "config/trial_design.json",
        "source/source_acquisition.csv",
        "references_verified.csv",
        "HOSTILE_REVIEW.md",
    ]:
        repro_entries.append((ROOT / filename, filename))
    for path in sorted((ROOT / "source" / "public").iterdir()):
        if path.is_file():
            repro_entries.append((path, str(path.relative_to(ROOT))))
    for path in sorted((ROOT / "literature" / "raw").iterdir()):
        if path.is_file():
            repro_entries.append((path, str(path.relative_to(ROOT))))
    repro_zip = ROOT / "reproducibility_bundle.zip"
    make_zip(repro_zip, repro_entries)
    results.append(("Reproducibility ZIP", f"{len(repro_entries)} intended files", "PASS"))

    qc_lines = [
        "# FINAL QC REPORT",
        "",
        f"Completed UTC: {NOW}",
        "",
        "| Check | Evidence | Result |",
        "|---|---|---|",
    ]
    qc_lines.extend(f"| {check} | {evidence} | {result} |" for check, evidence, result in results)
    qc_lines.extend(
        [
            "",
            "## Release blockers external to this package",
            "",
            "- ACA must confirm that it will consider a trial protocol and specify the portal article category.",
            "- Sponsor, funding, sites, author contributions, conflicts, insurance/compensation, ethics, registration, recruitment dates, and final eligibility thresholds must be confirmed.",
            "- The package is not authorised for enrolment and contains no prospective clinical results.",
        ]
    )
    report = "\n".join(qc_lines) + "\n"
    (ROOT / "FINAL_QC_REPORT.md").write_text(report, encoding="utf-8")
    (ROOT / "QC_REPORT.md").write_text(report, encoding="utf-8")

    submission_entries = [(ROOT / filename, filename) for filename in SUBMISSION_FILES]
    if standalone.exists():
        submission_entries.append((standalone, standalone.name))
    submission_zip = ROOT / "ACA_FPA_protocol_submission_package.zip"
    make_zip(submission_zip, submission_entries)
    results.append(("Submission ZIP", f"{len(submission_entries)} intended files", "PASS"))
    qc_lines = [
        "# FINAL QC REPORT",
        "",
        f"Completed UTC: {NOW}",
        "",
        "| Check | Evidence | Result |",
        "|---|---|---|",
    ]
    qc_lines.extend(f"| {check} | {evidence} | {result} |" for check, evidence, result in results)
    qc_lines.extend(
        [
            "",
            "## Release blockers external to this package",
            "",
            "- ACA must confirm that it will consider a trial protocol and specify the portal article category.",
            "- Sponsor, funding, sites, author contributions, conflicts, insurance/compensation, ethics, registration, recruitment dates, and final eligibility thresholds must be confirmed.",
            "- The package is not authorised for enrolment and contains no prospective clinical results.",
        ]
    )
    report = "\n".join(qc_lines) + "\n"
    (ROOT / "FINAL_QC_REPORT.md").write_text(report, encoding="utf-8")
    (ROOT / "QC_REPORT.md").write_text(report, encoding="utf-8")
    make_zip(submission_zip, submission_entries)

    manifest_rows = []
    for path in [repro_zip, submission_zip] + [ROOT / filename for filename in REQUIRED_DOCX]:
        manifest_rows.append(
            {
                "file": str(path.relative_to(ROOT)),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
        )
    with (ROOT / "output" / "final_file_manifest.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["file", "bytes", "sha256"],
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(manifest_rows)

    handoff = f"""FINAL HANDOFF

Canonical package: ACA_FPA_protocol_submission_package.zip
Reproducibility package: reproducibility_bundle.zip
Protocol: version 1.0 dated 2026-09-25
Design: multicentre, 1:1, double-blind, placebo-controlled; target N=800
Primary endpoint: sustained artefact-controlled relative MAP decline during 0–30 minutes after protamine start

REQUIRED BEFORE FORMAL SUBMISSION OR ENROLMENT
1. Send the ACA presubmission inquiry and obtain an article-type decision.
2. Replace all bracketed placeholders and confirm authorship, sponsor, funding, sites, ethics, registry, compensation, recruitment dates, and label-based eligibility thresholds.
3. Obtain final protocol/SAP signatures and all regulatory approvals.

The local retrospective PDF was generated from the supplied DOCX and is not independently verified as the final publisher PDF. Synthetic validation files are not clinical data.
"""
    (ROOT / "FINAL_HANDOFF.txt").write_text(handoff, encoding="utf-8")
    print(json.dumps({"checks": len(results), "submission_zip": str(submission_zip), "reproducibility_zip": str(repro_zip)}, indent=2))


def pd_read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


if __name__ == "__main__":
    main()
